"""tests/test_claude_artifacts.py

Integration verification for the Claude Artifacts workspace:
1. Router artifact transformation (no filler, no fences).
2. 1-Click quick actions & custom NLP transformations.
3. Version history progression (v1 -> v2 -> v3).
4. Rollback capability to previous versions.
"""

import unittest
from unittest.mock import MagicMock, patch
import streamlit as st

from core.llm_router import LLMRouter, transform_artifact
from core.ui_components import set_initial_artifact


class TestClaudeArtifactsWorkspace(unittest.TestCase):

    def setUp(self):
        # Clear Streamlit session state keys before each test
        for key in list(st.session_state.keys()):
            del st.session_state[key]
        self.router = LLMRouter()
        self.sample_text = (
            "Annual Revenue: $450 Million\n"
            "Operating Margin: 28.5%\n"
            "Net Profit: $128 Million\n"
            "Research Expenditure: $52 Million\n"
            "Risk Factors: High supply chain volatility observed in Q3."
        )

    def test_artifact_initialization(self):
        """Verify set_initial_artifact establishes v1 state with observations."""
        observations = {
            "Document Type": "TXT Document",
            "Page Count / Geometry": "1 Pages (50 Words)",
            "Key Entities": ["$450 Million", "28.5%"],
        }
        set_initial_artifact(
            title="Synthesized Document Analysis",
            content=self.sample_text,
            observations=observations
        )

        self.assertIn("artifact_versions", st.session_state)
        self.assertEqual(len(st.session_state.artifact_versions), 1)
        self.assertEqual(st.session_state.artifact_version_idx, 0)
        self.assertEqual(st.session_state.artifact_title, "Synthesized Document Analysis")
        self.assertEqual(st.session_state.artifact_observations["Document Type"], "TXT Document")

    def test_quick_action_convert_to_table(self):
        """Verify 'Convert to Table' produces clean Markdown table without preamble."""
        table_output = self.router.transform_artifact(
            current_content=self.sample_text,
            observations={"Document Type": "Financials"},
            user_instruction="Convert to Table",
            preferred_provider="offline"
        )

        self.assertIn("| Metric / Dimension | Specification & Context |", table_output)
        self.assertIn("| **Annual Revenue** | $450 Million |", table_output)
        self.assertIn("| **Operating Margin** | 28.5% |", table_output)
        self.assertNotIn("```markdown", table_output)
        self.assertNotIn("```", table_output)
        self.assertNotIn("Sure!", table_output)

    def test_quick_action_executive_summary(self):
        """Verify 'Executive Summary' produces clean executive brief."""
        exec_output = self.router.transform_artifact(
            current_content=self.sample_text,
            observations={"Document Type": "Financials"},
            user_instruction="Executive Summary",
            preferred_provider="offline"
        )

        self.assertIn("# ⚡ Executive Summary", exec_output)
        self.assertIn("Key Findings", exec_output)
        self.assertNotIn("```markdown", exec_output)

    def test_custom_prompt_translation(self):
        """Verify custom NLP instruction (Translate to French) executes."""
        french_output = self.router.transform_artifact(
            current_content=self.sample_text,
            observations={"Document Type": "Financials"},
            user_instruction="Translate to French",
            preferred_provider="offline"
        )

        self.assertIn("Version Française", french_output)
        self.assertNotIn("```", french_output)

    def test_version_progression_and_rollback(self):
        """Verify v1 -> v2 -> v3 version progression and rollback to v1."""
        # 1. Start with v1
        set_initial_artifact(title="Test Artifact", content=self.sample_text)
        self.assertEqual(len(st.session_state.artifact_versions), 1)
        self.assertEqual(st.session_state.artifact_version_idx, 0)

        # 2. Add v2 (Convert to Table)
        v2_text = self.router.transform_artifact(
            current_content=st.session_state.artifact_versions[st.session_state.artifact_version_idx],
            observations={},
            user_instruction="Convert to Table",
            preferred_provider="offline"
        )
        st.session_state.artifact_versions.append(v2_text)
        st.session_state.artifact_version_idx = len(st.session_state.artifact_versions) - 1
        self.assertEqual(len(st.session_state.artifact_versions), 2)
        self.assertEqual(st.session_state.artifact_version_idx, 1)
        self.assertIn("| Metric / Dimension |", st.session_state.artifact_versions[1])

        # 3. Add v3 (Translate to French)
        v3_text = self.router.transform_artifact(
            current_content=st.session_state.artifact_versions[st.session_state.artifact_version_idx],
            observations={},
            user_instruction="Translate to French",
            preferred_provider="offline"
        )
        st.session_state.artifact_versions.append(v3_text)
        st.session_state.artifact_version_idx = len(st.session_state.artifact_versions) - 1
        self.assertEqual(len(st.session_state.artifact_versions), 3)
        self.assertEqual(st.session_state.artifact_version_idx, 2)
        self.assertIn("Version Française", st.session_state.artifact_versions[2])

        # 4. Rollback to v1
        st.session_state.artifact_version_idx = 0
        self.assertEqual(st.session_state.artifact_version_idx, 0)
        self.assertEqual(st.session_state.artifact_versions[st.session_state.artifact_version_idx], self.sample_text)

    @patch("core.llm_router.query_llm")
    def test_cloud_llm_fence_and_filler_stripping(self, mock_query):
        """Verify cloud LLM conversational fillers and code fences are completely stripped."""
        mock_query.return_value = (
            "Here is the updated document you requested:\n\n"
            "```markdown\n"
            "# Strategic Plan 2026\n\n"
            "- Initiative 1: Cloud Expansion\n"
            "- Initiative 2: AI Automation\n"
            "```\n\n"
            "Let me know if you need anything else!"
        )

        res = transform_artifact(
            current_content="Old plan",
            observations={},
            user_instruction="Update plan",
            preferred_provider="gemini"
        )

        self.assertTrue(res.startswith("# Strategic Plan 2026"))
        self.assertNotIn("```markdown", res)
        self.assertNotIn("Here is the updated document", res)
        self.assertNotIn("Let me know if you need anything else", res)

    def test_export_markdown_to_docx(self):
        """Verify export_markdown_to_docx generates in-memory docx with tables and headers."""
        from core.ui_components import export_markdown_to_docx
        from docx import Document

        sample_md = (
            "# Executive Brief\n\n"
            "This is a paragraph with **bold metrics** and `code` tags.\n\n"
            "| Revenue | Growth |\n"
            "| :--- | :--- |\n"
            "| $450M | 28% |\n\n"
            "- Strategic priority 1\n"
            "- Strategic priority 2\n\n"
            "> Approved by audit committee."
        )

        buf = export_markdown_to_docx(sample_md, title="Executive Report")
        self.assertGreater(len(buf.getvalue()), 1000)

        # Re-parse generated document to verify structure
        parsed_doc = Document(buf)
        self.assertGreater(len(parsed_doc.paragraphs), 0)
        self.assertEqual(len(parsed_doc.tables), 1)
        self.assertEqual(len(parsed_doc.tables[0].rows), 2)
        self.assertEqual(parsed_doc.tables[0].cell(0, 0).text, "Revenue")
        self.assertEqual(parsed_doc.tables[0].cell(1, 0).text, "$450M")

    def test_export_markdown_to_pdf(self):
        """Verify export_markdown_to_pdf generates valid in-memory PDF bytes with tables."""
        from core.ui_components import export_markdown_to_pdf

        sample_md = (
            "# Executive Brief\n\n"
            "This is a paragraph with key findings.\n\n"
            "| Metric | Value |\n"
            "| :--- | :--- |\n"
            "| Revenue | $450M |\n\n"
            "- Operational metric\n"
            "> Confirmed by executive oversight."
        )

        buf = export_markdown_to_pdf(sample_md, title="Executive Report")
        pdf_bytes = buf.getvalue()
        self.assertGreater(len(pdf_bytes), 500)
        # PDF files always start with %PDF header
        self.assertTrue(pdf_bytes.startswith(b"%PDF"))

    def test_compute_markdown_diff_html(self):
        """Verify compute_markdown_diff_html generates safe, styled diff markup."""
        from core.ui_components import compute_markdown_diff_html

        old_text = "Line 1: Initial\nLine 2: Delete Me\nLine 3: Keep"
        new_text = "Line 1: Initial\nLine 2: Added New\nLine 3: Keep"

        diff_html = compute_markdown_diff_html(old_text, new_text)

        # Verify additions styling
        self.assertIn("#30D158", diff_html)
        self.assertIn("Line 2: Added New", diff_html)
        self.assertIn("diff-add", diff_html)

        # Verify deletions styling & strikethrough
        self.assertIn("#FF453A", diff_html)
        self.assertIn("line-through", diff_html)
        self.assertIn("Line 2: Delete Me", diff_html)
        self.assertIn("diff-del", diff_html)

        # Verify chunk header styling
        self.assertIn("#FFAA55", diff_html)
        self.assertIn("diff-chunk", diff_html)

        # Verify unchanged text
        self.assertIn("#A1A1A6", diff_html)
        self.assertIn("Line 1: Initial", diff_html)

        # Verify scrollable container
        self.assertIn("diff-inspector-container", diff_html)
        self.assertIn("max-height: 520px", diff_html)

    def test_compute_markdown_diff_identical(self):
        """Verify identical versions produce clean informative notice."""
        from core.ui_components import compute_markdown_diff_html

        text = "# Section Header\nContent line 1\nContent line 2"
        diff_html = compute_markdown_diff_html(text, text)
        self.assertIn("Identical Versions", diff_html)
        self.assertIn("diff-inspector-container", diff_html)

    def test_compute_markdown_diff_html_escaping(self):
        """Verify HTML tags in diffed markdown are safely escaped."""
        from core.ui_components import compute_markdown_diff_html

        old_text = "<script>alert('xss')</script>"
        new_text = "<script>alert('safe')</script>"
        diff_html = compute_markdown_diff_html(old_text, new_text)
        self.assertNotIn("<script>", diff_html)
        self.assertIn("&lt;script&gt;", diff_html)


if __name__ == "__main__":
    unittest.main()

